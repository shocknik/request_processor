"""Производительность и изоляция extraction от ассистента (.docx)."""

from __future__ import annotations

from pathlib import Path
from unittest.mock import patch

import pytest

from request_processor.extraction.pdf_extractor import (
    _clean_mark,
    _collapse_horizontal_merge_cells,
    _compact_text_for_marks,
    _dedupe_consecutive_lines,
    extract_from_document,
    find_cable_marks,
    load_docx_content,
)

_TRAINING_DOCX = (
    Path(__file__).resolve().parents[1]
    / "data"
    / "training"
    / "documents"
    / "registered"
    / "11.02.25.2 Направление в ИЛ 10067033 Кабель-тест.docx"
)


def test_collapse_horizontal_merge_cells() -> None:
    raw = ["", "A", "A", "A", "B", "B", ""]
    assert _collapse_horizontal_merge_cells(raw) == ["", "A", "B", ""]


def test_dedupe_consecutive_lines() -> None:
    text = "A\nA\nA\nB\nB\n\n\nC"
    assert _dedupe_consecutive_lines(text) == "A\nB\n\nC"


def test_compact_text_for_marks_shrinks_repeats() -> None:
    bloated = "\n".join(["В испытательную лабораторию (ИЛ)"] * 200)
    bloated += "\nМарка МКУПнг(А)-LS 2х2x0,35\n"
    compact = _compact_text_for_marks(bloated)
    assert compact.count("В испытательную") == 1
    assert "МКУПнг" in compact
    assert len(compact) < len(bloated)


def test_clean_mark_is_deterministic_no_assistant() -> None:
    """_clean_mark не должен трогать MarkCorrector / SQLite / Ollama."""
    with patch(
        "request_processor.assistant.mark_corrector.suggest_mark_correction",
        side_effect=AssertionError("suggest_mark_correction must not be called"),
    ):
        with patch(
            "request_processor.assistant.mark_corrector.get_mark_corrector",
            side_effect=AssertionError("get_mark_corrector must not be called"),
        ):
            cleaned = _clean_mark("  МКУПнг(A)-LS 2х2x0,35  ")
    assert "МКУП" in cleaned or "МКУПнг" in cleaned
    assert cleaned == cleaned.strip()


def test_find_cable_marks_no_assistant_on_repeats() -> None:
    """Повторы одной марки → мало уникальных; ассистент не вызывается."""
    line = "МКУПнг(А)-LS 2х2x0,35"
    text = "\n".join([f"Продукция: {line}"] * 500)
    with patch(
        "request_processor.assistant.mark_corrector.suggest_mark_correction",
        side_effect=AssertionError("assistant must not run in find_cable_marks"),
    ):
        with patch(
            "request_processor.assistant.mark_corrector.MarkCorrector.suggest",
            side_effect=AssertionError("MarkCorrector.suggest must not run"),
        ):
            marks = find_cable_marks(text)
    assert len(marks) >= 1
    assert len({m.mark.lower() for m in marks}) <= 3


@pytest.mark.skipif(not _TRAINING_DOCX.is_file(), reason="training docx отсутствует")
def test_load_docx_content_single_pass_not_bloated() -> None:
    text, tables = load_docx_content(_TRAINING_DOCX)
    flat = "\n".join(" | ".join(c for c in row if c) for t in tables for row in t)
    assert len(tables) >= 1
    assert len(flat) < 25_000, f"tables still bloated: {len(flat)} chars"
    assert "Кабель-Тест" in flat or "Кабель" in flat or "МКУП" in flat + text


@pytest.mark.skipif(not _TRAINING_DOCX.is_file(), reason="training docx отсутствует")
def test_extract_docx_opens_document_once() -> None:
    """extract_from_document(.docx) открывает python-docx Document ровно один раз."""
    from docx import Document as RealDocument

    call_count = {"n": 0}

    def counting_document(*args, **kwargs):
        call_count["n"] += 1
        return RealDocument(*args, **kwargs)

    with patch("docx.Document", side_effect=counting_document):
        result = extract_from_document(_TRAINING_DOCX, use_ocr=False)

    assert result.source_type == "docx"
    assert len(result.cable_marks) >= 1
    assert call_count["n"] == 1, f"Document opened {call_count['n']} times"


@pytest.mark.skipif(not _TRAINING_DOCX.is_file(), reason="training docx отсутствует")
def test_extract_docx_does_not_call_assistant() -> None:
    """Базовый parser path не вызывает suggest_mark_correction / MarkCorrector."""
    with patch(
        "request_processor.assistant.mark_corrector.suggest_mark_correction",
        side_effect=AssertionError("suggest must not run during extract"),
    ):
        with patch(
            "request_processor.assistant.mark_corrector.MarkCorrector.suggest",
            side_effect=AssertionError("MarkCorrector.suggest must not run"),
        ):
            result = extract_from_document(_TRAINING_DOCX, use_ocr=False)
    assert result.source_type == "docx"
    assert len(result.cable_marks) >= 1


@pytest.mark.skipif(not _TRAINING_DOCX.is_file(), reason="training docx отсутствует")
def test_extract_docx_fast_and_finds_marks() -> None:
    import time

    t0 = time.perf_counter()
    result = extract_from_document(_TRAINING_DOCX, use_ocr=False)
    elapsed = time.perf_counter() - t0
    assert result.source_type == "docx"
    assert len(result.cable_marks) >= 1, result.cable_marks
    # после отделения ассистента — цель < 2 с (CI + cold import)
    assert elapsed < 2.0, f"docx extract too slow: {elapsed:.2f}s"
    marks_blob = " ".join(m.mark for m in result.cable_marks)
    assert "МКУП" in marks_blob or "нг" in marks_blob.lower()


def test_suggest_many_dedupes_identical_marks(tmp_path: Path) -> None:
    """Пакетный corrector: одинаковые строки → один fuzzy, не N."""
    from request_processor.assistant.mark_corrector import MarkCorrector
    from request_processor.persistence.sqlite_repo import init_db

    db = tmp_path / "assist.db"
    init_db(db)
    corrector = MarkCorrector(db_path=db)
    n_calls = {"n": 0}
    orig = corrector._suggest_no_llm

    def counting(raw, *, context):
        n_calls["n"] += 1
        return orig(raw, context=context)

    corrector._suggest_no_llm = counting  # type: ignore[method-assign]
    marks = ["МКУПнг(А)-LS 2х2x0,35"] * 20 + ["ВВГнг(А) 3х1,5"] * 10
    out = corrector.suggest_many(marks, only_changed=False, use_llm=False)
    assert n_calls["n"] == 2, f"expected 2 unique computations, got {n_calls['n']}"
    assert len(out) == 30
